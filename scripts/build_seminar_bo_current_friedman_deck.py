from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SLIDES_DIR = ROOT / "slides"
TEMPLATE_DECK = SLIDES_DIR / "260412_A.pptx"
OUTPUT_PPTX = SLIDES_DIR / "seminar_bo_current_friedman.pptx"
RESULT_DIR = (
    ROOT
    / "outputs"
    / "seminar_bo_current"
    / "sr_alpha_friedman"
    / "seminar_bo_current_seed100_20260602"
)


COLOR_TITLE = RGBColor(0x17, 0x32, 0x4D)
COLOR_BODY = RGBColor(0x31, 0x47, 0x5E)
COLOR_MUTED = RGBColor(0x68, 0x76, 0x86)
COLOR_BLUE = RGBColor(0x2F, 0x6B, 0x9A)
COLOR_GREEN = RGBColor(0x2E, 0x8B, 0x57)
COLOR_AMBER = RGBColor(0xC7, 0x7D, 0x1A)
COLOR_RED = RGBColor(0xC4, 0x45, 0x36)
COLOR_PURPLE = RGBColor(0x7A, 0x4E, 0x76)
COLOR_LINE = RGBColor(0xD8, 0xE0, 0xE8)

FILL_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FILL_GRAY = RGBColor(0xF3, 0xF6, 0xF9)
FILL_BLUE = RGBColor(0xEA, 0xF3, 0xFF)
FILL_GREEN = RGBColor(0xEA, 0xF7, 0xEF)
FILL_AMBER = RGBColor(0xFF, 0xF4, 0xE5)
FILL_RED = RGBColor(0xFD, 0xEC, 0xEB)
FILL_PURPLE = RGBColor(0xF3, 0xEC, 0xF5)


def remove_all_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[3])


def add_footer(slide, index: int) -> None:
    add_text(
        slide,
        0.62,
        6.96,
        10.4,
        0.24,
        "BO-controlled GP / Seminar result / Friedman-I",
        size=8.5,
        color=COLOR_MUTED,
    )
    add_text(
        slide,
        12.35,
        6.96,
        0.7,
        0.24,
        str(index),
        size=8.5,
        color=COLOR_MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(slide, 0.62, 0.34, 11.9, 0.42, title, size=24, bold=True, color=COLOR_TITLE)
    if subtitle:
        add_text(slide, 0.64, 0.82, 11.1, 0.3, subtitle, size=12.5, color=COLOR_MUTED)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.62),
        Inches(1.03),
        Inches(11.7),
        Inches(0.025),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_LINE
    line.line.fill.background()


def add_text(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: float = 16,
    color: RGBColor = COLOR_BODY,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_paragraphs(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    paragraphs: list[dict],
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    rounded: bool = False,
):
    if fill is not None or line is not None or rounded:
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        if fill is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        else:
            shape.fill.background()
        if line is not None:
            shape.line.color.rgb = line
            shape.line.width = Pt(1.2)
        else:
            shape.line.fill.background()
    else:
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))

    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    for idx, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item["text"]
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(item.get("space_after", 3))
        p.level = item.get("level", 0)
        for run in p.runs:
            run.font.size = Pt(item.get("size", 15))
            run.font.bold = item.get("bold", False)
            run.font.color.rgb = item.get("color", COLOR_BODY)
    return shape


def add_panel(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
    lines: list[str],
    fill: RGBColor,
    line: RGBColor,
    title_size: float = 17,
    body_size: float = 13.5,
    center: bool = False,
):
    paragraphs = [
        {
            "text": title,
            "size": title_size,
            "bold": True,
            "color": COLOR_TITLE,
            "align": PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
            "space_after": 5,
        }
    ]
    for text in lines:
        paragraphs.append(
            {
                "text": text,
                "size": body_size,
                "color": COLOR_BODY,
                "align": PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                "space_after": 2,
            }
        )
    return add_paragraphs(slide, left, top, width, height, paragraphs, fill=fill, line=line, rounded=True)


def add_metric_card(
    slide,
    left: float,
    top: float,
    width: float,
    title: str,
    value: str,
    note: str,
    fill: RGBColor,
    line: RGBColor,
):
    return add_paragraphs(
        slide,
        left,
        top,
        width,
        1.24,
        [
            {"text": title, "size": 12.5, "bold": True, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER},
            {"text": value, "size": 24, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER},
            {"text": note, "size": 10.5, "color": COLOR_BODY, "align": PP_ALIGN.CENTER},
        ],
        fill=fill,
        line=line,
        rounded=True,
    )


def add_picture_contain(
    slide,
    image_path: Path,
    left: float,
    top: float,
    max_width: float,
    max_height: float,
    center_x: bool = True,
    center_y: bool = True,
):
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    scale = min(max_width / image_width, max_height / image_height)
    width = image_width * scale
    height = image_height * scale
    adjusted_left = left + (max_width - width) / 2 if center_x else left
    adjusted_top = top + (max_height - height) / 2 if center_y else top
    return slide.shapes.add_picture(
        str(image_path),
        Inches(adjusted_left),
        Inches(adjusted_top),
        width=Inches(width),
        height=Inches(height),
    )


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = COLOR_MUTED):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(2.0)
    line.line.end_arrowhead = True
    return line


def slide_1(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_text(
        slide,
        0.9,
        1.35,
        11.4,
        0.82,
        "文脈付き BO による\n多目的 GP の操作率・更新周期制御",
        size=30,
        bold=True,
        color=COLOR_TITLE,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        2.0,
        3.05,
        9.2,
        0.42,
        "Seminar 用本実験結果: Friedman-I symbolic regression",
        size=18,
        color=COLOR_BODY,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        3.1,
        4.2,
        7.0,
        0.34,
        "bo_current 方針での中間評価",
        size=19,
        bold=True,
        color=COLOR_RED,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, 4.2, 5.55, 4.8, 0.3, "2026年6月2日", size=14, color=COLOR_MUTED, align=PP_ALIGN.CENTER)


def slide_2(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "今日のメッセージ", "Seminar では「動いたこと」と「まだ足りないこと」を両方示す")
    add_panel(
        slide,
        0.85,
        1.45,
        3.75,
        3.85,
        "1. 閉ループ制御は動作",
        ["GP 状態を観測し、BO が p_c, p_m, k を逐次選択できた"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=18,
        body_size=15.5,
    )
    add_panel(
        slide,
        4.85,
        1.45,
        3.75,
        3.85,
        "2. 標準固定率は上回る",
        ["Friedman-I で standard 固定率より HV 平均が高い"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=18,
        body_size=15.5,
    )
    add_panel(
        slide,
        8.85,
        1.45,
        3.75,
        3.85,
        "3. 最良固定率には未達",
        ["高突然変異固定率が最良。報酬設計と BO 制御器に改善余地"],
        FILL_AMBER,
        COLOR_AMBER,
        title_size=18,
        body_size=15.5,
    )
    add_text(
        slide,
        1.15,
        5.92,
        10.9,
        0.42,
        "結論: 提案枠組みの妥当性は確認できたが、現行 bo_current は改善途中である。",
        size=18,
        bold=True,
        color=COLOR_TITLE,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 2)


def slide_3(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "研究背景: 固定率 GP の限界", "操作率は探索性・収束性・停滞打破に強く関わる")
    add_panel(
        slide,
        0.8,
        1.45,
        3.45,
        3.5,
        "固定率 GP",
        ["p_c, p_m を全世代で固定", "実装は簡単", "ただし問題・世代依存性を扱いにくい"],
        FILL_GRAY,
        COLOR_LINE,
        title_size=18,
    )
    add_panel(
        slide,
        4.72,
        1.45,
        3.45,
        3.5,
        "探索状態の変化",
        ["初期: 多様性を確保", "中盤: 有望解を組み替え", "停滞時: 再探索が必要"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=18,
    )
    add_panel(
        slide,
        8.65,
        1.45,
        3.45,
        3.5,
        "本研究の立場",
        ["操作率を制御入力とみなす", "GP をプラントとして観測", "BO を制御器として使う"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=18,
    )
    add_footer(slide, 3)


def slide_4(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "提案法 bo_current の閉ループ構造", "GP 状態を観測し、p_c, p_m, k を文脈付き BO が選ぶ")
    y = 2.25
    boxes = [
        (0.85, "State\nx_l", FILL_BLUE, COLOR_BLUE),
        (3.35, "Contextual BO\nper-k EI", FILL_GREEN, COLOR_GREEN),
        (6.05, "Action\n(p_c, p_m, k)", FILL_AMBER, COLOR_AMBER),
        (8.75, "MO-GP plant\nk generations", FILL_GRAY, COLOR_LINE),
    ]
    for left, label, fill, line in boxes:
        add_panel(slide, left, y, 2.05, 1.25, label, [], fill, line, title_size=17, center=True)
    add_arrow(slide, 2.9, y + 0.62, 3.35, y + 0.62)
    add_arrow(slide, 5.4, y + 0.62, 6.05, y + 0.62)
    add_arrow(slide, 8.1, y + 0.62, 8.75, y + 0.62)
    add_arrow(slide, 10.0, 3.6, 1.55, 4.7, COLOR_RED)
    add_text(slide, 1.7, 4.72, 8.6, 0.32, "区間評価: ΔHV_rate, 平均多様性, 制御コスト → 報酬 r_l", size=15.5, color=COLOR_RED, bold=True, align=PP_ALIGN.CENTER)
    add_panel(
        slide,
        1.1,
        5.35,
        10.9,
        0.85,
        "今回の本線",
        ["EI の基準値は k ごとの過去最高報酬を使う per-k 方式。beta_EI は Seminar 後に再検討。"],
        FILL_RED,
        COLOR_RED,
        title_size=15,
        body_size=13.8,
        center=True,
    )
    add_footer(slide, 4)


def slide_5(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "bo_current の数理的な整理", "状態・行動・報酬・EI の定義")
    add_panel(
        slide,
        0.75,
        1.35,
        3.8,
        2.05,
        "状態 x_l",
        ["世代進行率", "現在 HV / 直近改善", "多様性 / 木サイズ / 停滞"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=17,
    )
    add_panel(
        slide,
        4.85,
        1.35,
        3.8,
        2.05,
        "行動 u_l",
        ["交叉率 p_c", "突然変異率 p_m", "更新周期 k ∈ {1,3,5}"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=17,
    )
    add_panel(
        slide,
        8.95,
        1.35,
        3.2,
        2.05,
        "報酬 r_l",
        ["HV 改善率を主項", "多様性維持を副項", "制御コストを抑制"],
        FILL_AMBER,
        COLOR_AMBER,
        title_size=17,
    )
    add_paragraphs(
        slide,
        1.05,
        4.08,
        11.15,
        1.35,
        [
            {"text": "bo_current の EI", "size": 17.5, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER},
            {"text": "r_best^(k) = max { r_i | k_i = k }", "size": 21, "bold": True, "color": COLOR_RED, "align": PP_ALIGN.CENTER},
            {"text": "各 k の中で「今までより良くなりそうか」を評価する", "size": 14.5, "color": COLOR_BODY, "align": PP_ALIGN.CENTER},
        ],
        fill=FILL_GRAY,
        line=COLOR_LINE,
        rounded=True,
    )
    add_footer(slide, 5)


def slide_6(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "対象問題: Friedman-I symbolic regression", "精度と式の簡潔さの Pareto front を求める")
    add_paragraphs(
        slide,
        0.88,
        1.35,
        6.2,
        1.45,
        [
            {"text": "目標関数", "size": 16, "bold": True, "color": COLOR_TITLE},
            {"text": "y = 10 sin(pi x1 x2) + 20(x3 - 0.5)^2 + 10x4 + 5x5", "size": 17.5, "bold": True, "color": COLOR_BODY},
            {"text": "x_i は U(0,1) から生成。訓練点 200、テスト点 200。", "size": 13.5, "color": COLOR_MUTED},
        ],
        fill=FILL_BLUE,
        line=COLOR_BLUE,
        rounded=True,
    )
    add_panel(
        slide,
        7.55,
        1.35,
        4.45,
        1.45,
        "2目的最適化",
        ["目的1: training NRMSE を最小化", "目的2: tree size を最小化"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=16,
        body_size=14.5,
    )
    add_panel(
        slide,
        0.88,
        3.35,
        3.65,
        1.72,
        "なぜこの問題か",
        ["GP の代表的な回帰 benchmark", "非線形・相互作用を含む"],
        FILL_GRAY,
        COLOR_LINE,
        title_size=16,
    )
    add_panel(
        slide,
        4.9,
        3.35,
        3.65,
        1.72,
        "評価しやすい点",
        ["HV で Pareto front を比較", "式サイズとの trade-off が明確"],
        FILL_AMBER,
        COLOR_AMBER,
        title_size=16,
    )
    add_panel(
        slide,
        8.92,
        3.35,
        3.1,
        1.72,
        "今回の役割",
        ["Seminar 用の主対象問題", "bo_current の妥当性確認"],
        FILL_RED,
        COLOR_RED,
        title_size=16,
    )
    add_footer(slide, 6)


def slide_7(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "実験条件", "固定率 GP 3条件と bo_current を 100 seed で比較")
    add_metric_card(slide, 0.82, 1.3, 2.6, "seed", "100", "0..99", FILL_BLUE, COLOR_BLUE)
    add_metric_card(slide, 3.68, 1.3, 2.6, "population", "24", "individuals", FILL_GREEN, COLOR_GREEN)
    add_metric_card(slide, 6.54, 1.3, 2.6, "generations", "30", "720 evaluations/run", FILL_AMBER, COLOR_AMBER)
    add_metric_card(slide, 9.4, 1.3, 2.6, "archive key", "topo+val", "structure + values", FILL_GRAY, COLOR_LINE)
    methods = [
        ("standard", "p_c=0.80, p_m=0.05", FILL_GRAY, COLOR_LINE),
        ("high mutation", "p_c=0.70, p_m=0.20", FILL_GREEN, COLOR_GREEN),
        ("high crossover", "p_c=0.90, p_m=0.05", FILL_AMBER, COLOR_AMBER),
        ("bo_current", "p_c,p_m,k を BO で制御", FILL_RED, COLOR_RED),
    ]
    for i, (title, body, fill, line) in enumerate(methods):
        add_panel(slide, 0.9 + i * 3.05, 3.55, 2.65, 1.45, title, [body], fill, line, title_size=16, body_size=13.2, center=True)
    add_text(slide, 1.4, 5.85, 10.3, 0.34, "比較の焦点: 固定率と比べて、状態依存制御が HV と多様性をどこまで改善できるか", size=16, bold=True, color=COLOR_TITLE, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)


def slide_8(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "結果1: final HV と diversity", "bo_current は standard と high crossover を上回るが、high mutation には僅差で未達")
    add_picture_contain(slide, RESULT_DIR / "seminar_final_hv_diversity.png", 0.72, 1.2, 8.2, 5.25)
    add_panel(
        slide,
        9.25,
        1.55,
        3.0,
        1.18,
        "HV平均",
        ["bo_current: 0.7731", "standard: 0.7570", "high mutation: 0.7766"],
        FILL_RED,
        COLOR_RED,
        title_size=15,
        body_size=12.5,
    )
    add_panel(
        slide,
        9.25,
        3.05,
        3.0,
        1.18,
        "読み取り",
        ["標準固定率は上回る", "最良固定率には届かない"],
        FILL_AMBER,
        COLOR_AMBER,
        title_size=15,
        body_size=12.5,
    )
    add_panel(
        slide,
        9.25,
        4.55,
        3.0,
        1.18,
        "発表での主張",
        ["現行版は有望だが", "改善途中である"],
        FILL_GRAY,
        COLOR_LINE,
        title_size=15,
        body_size=12.5,
    )
    add_footer(slide, 8)


def slide_9(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "結果2: seed 内差分で見る", "平均だけでなく、同じ seed 内で bo_current と固定率を比較する")
    add_picture_contain(slide, RESULT_DIR / "seminar_paired_hv_difference.png", 0.72, 1.2, 6.65, 4.95)
    add_paragraphs(
        slide,
        7.65,
        1.35,
        4.6,
        4.35,
        [
            {"text": "seed 内 final HV 差分", "size": 17, "bold": True, "color": COLOR_TITLE},
            {"text": "vs standard: +0.0161 / win 64", "size": 15, "color": COLOR_GREEN, "bold": True},
            {"text": "vs high mutation: -0.0035 / win 40", "size": 15, "color": COLOR_RED, "bold": True},
            {"text": "vs high crossover: +0.0084 / win 54", "size": 15, "color": COLOR_GREEN, "bold": True},
            {"text": "vs best fixed per seed: -0.0201 / win 21", "size": 15, "color": COLOR_RED, "bold": True},
            {"text": "結論: standard baseline には優位傾向。ただし tuned baseline を安定して上回るには未達。", "size": 13, "color": COLOR_BODY},
        ],
        fill=FILL_GRAY,
        line=COLOR_LINE,
        rounded=True,
    )
    add_footer(slide, 9)


def slide_10(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "結果3: HV 分布と外れ値", "bo_current の中心は standard より高いが、外れ値はまだ残る")
    add_picture_contain(slide, RESULT_DIR / "seminar_final_hv_boxplot.png", 0.82, 1.25, 7.3, 4.85)
    add_panel(
        slide,
        8.55,
        1.55,
        3.45,
        1.42,
        "分布の見方",
        ["中央値は各手法で大差ない", "低HV側の外れ値が性能を下げる"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=15.5,
        body_size=12.8,
    )
    add_panel(
        slide,
        8.55,
        3.25,
        3.45,
        1.42,
        "bo_current の課題",
        ["悪い seed を減らす制御が必要", "報酬設計と探索周期の調整が候補"],
        FILL_RED,
        COLOR_RED,
        title_size=15.5,
        body_size=12.8,
    )
    add_footer(slide, 10)


def slide_11(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "結果4: bo_current の k 選択傾向", "毎世代更新ではなく、平均 11.26 回の制御更新で進化を進めた")
    add_picture_contain(slide, RESULT_DIR / "seminar_k_selection_counts.png", 0.82, 1.35, 5.4, 4.6)
    add_metric_card(slide, 6.85, 1.45, 2.65, "mean p_c", "0.7025", "交叉率の平均", FILL_BLUE, COLOR_BLUE)
    add_metric_card(slide, 9.75, 1.45, 2.65, "mean p_m", "0.1381", "突然変異率の平均", FILL_GREEN, COLOR_GREEN)
    add_metric_card(slide, 6.85, 3.15, 2.65, "control steps", "11.26", "30世代中の平均更新回数", FILL_AMBER, COLOR_AMBER)
    add_metric_card(slide, 9.75, 3.15, 2.65, "k choices", "491/333/302", "k=1 / 3 / 5", FILL_RED, COLOR_RED)
    add_text(slide, 6.85, 5.3, 5.45, 0.52, "k=1 が最多だが、k=3,5 も使われており、更新周期を制御対象にした意味は確認できる。", size=14.2, bold=True, color=COLOR_TITLE, align=PP_ALIGN.CENTER)
    add_footer(slide, 11)


def slide_12(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "考察: 何が分かったか", "勝ち負けより、現行版の性質が見えたことが今回の収穫")
    add_panel(
        slide,
        0.88,
        1.35,
        3.65,
        3.8,
        "良かった点",
        ["標準固定率より HV 平均が高い", "制御周期 k を実際に切替", "閉ループ実験基盤が整った"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=17,
        body_size=14,
    )
    add_panel(
        slide,
        4.82,
        1.35,
        3.65,
        3.8,
        "足りない点",
        ["高突然変異固定率が最良", "best fixed per seed には未達", "低HV seed を抑えきれない"],
        FILL_RED,
        COLOR_RED,
        title_size=17,
        body_size=14,
    )
    add_panel(
        slide,
        8.76,
        1.35,
        3.65,
        3.8,
        "解釈",
        ["Friedman-I では探索性が重要", "現行報酬は突然変異寄りの強い固定率を十分に再現できていない可能性"],
        FILL_AMBER,
        COLOR_AMBER,
        title_size=17,
        body_size=14,
    )
    add_text(slide, 1.2, 5.76, 11.0, 0.35, "Seminar では「bo_current は有望だが、次の改善が必要」という結論で進める。", size=16.5, bold=True, color=COLOR_TITLE, align=PP_ALIGN.CENTER)
    add_footer(slide, 12)


def slide_13(prs: Presentation) -> None:
    slide = blank_slide(prs)
    add_title(slide, "今後の方針", "Seminar 後に beta_EI と k 周りを再検討する")
    add_panel(
        slide,
        0.92,
        1.35,
        3.65,
        3.65,
        "1. bo_current の改善",
        ["報酬内訳の分析", "低HV seed の原因確認", "多様性項・制御コストの調整"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=17,
        body_size=14,
    )
    add_panel(
        slide,
        4.85,
        1.35,
        3.65,
        3.65,
        "2. 比較設計の強化",
        ["固定率グリッドの拡張", "統計検定の追加", "他問題への展開"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=17,
        body_size=14,
    )
    add_panel(
        slide,
        8.78,
        1.35,
        3.65,
        3.65,
        "3. Seminar 後検討",
        ["beta_EI の再評価", "k の warm-up / EI 比較", "BO surrogate の安定化"],
        FILL_PURPLE,
        COLOR_PURPLE,
        title_size=17,
        body_size=14,
    )
    add_text(slide, 1.25, 5.75, 10.7, 0.38, "今回の発表では bo_current を本線とし、改良案は今後の課題として整理する。", size=17, bold=True, color=COLOR_TITLE, align=PP_ALIGN.CENTER)
    add_footer(slide, 13)


def build_deck() -> Path:
    prs = Presentation(str(TEMPLATE_DECK))
    remove_all_slides(prs)
    prs.core_properties.title = "seminar_bo_current_friedman"
    prs.core_properties.subject = "Seminar result deck for BO-controlled multi-objective GP"

    slide_1(prs)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs)
    slide_6(prs)
    slide_7(prs)
    slide_8(prs)
    slide_9(prs)
    slide_10(prs)
    slide_11(prs)
    slide_12(prs)
    slide_13(prs)

    prs.save(str(OUTPUT_PPTX))
    return OUTPUT_PPTX


if __name__ == "__main__":
    print(build_deck())
