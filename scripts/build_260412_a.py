from __future__ import annotations

from pathlib import Path
import subprocess
import tempfile
import time

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SLIDES_DIR = ROOT / "slides"
TEMPLATE_PATH = SLIDES_DIR / "251022_明神.pptx"
OUTPUT_PATH = SLIDES_DIR / "260412_A.pptx"
GENERATED_DIR = SLIDES_DIR / "generated_assets" / "260412_A"
CHROME_PATH = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")

FIGURE_MAP = {
    "slide07": SLIDES_DIR / "figures" / "slide07_closed_loop_architecture.svg",
    "slide08": SLIDES_DIR / "figures" / "slide08_control_step_flow.svg",
    "slide10": SLIDES_DIR / "figures" / "slide10_context_vector_design.svg",
    "slide12": SLIDES_DIR / "figures" / "slide12_reward_function_design.svg",
    "slide15": SLIDES_DIR / "figures" / "slide15_contextual_bo_reason.svg",
}


COLOR_TITLE = RGBColor(0x17, 0x32, 0x4D)
COLOR_BODY = RGBColor(0x31, 0x47, 0x5E)
COLOR_MUTED = RGBColor(0x5F, 0x6F, 0x82)
COLOR_BLUE = RGBColor(0x2F, 0x6B, 0x9A)
COLOR_GREEN = RGBColor(0x2E, 0x8B, 0x57)
COLOR_AMBER = RGBColor(0xC7, 0x7D, 0x1A)
COLOR_RED = RGBColor(0xC4, 0x45, 0x36)

FILL_BLUE = RGBColor(0xEA, 0xF3, 0xFF)
FILL_GREEN = RGBColor(0xEA, 0xF7, 0xEF)
FILL_AMBER = RGBColor(0xFF, 0xF4, 0xE5)
FILL_RED = RGBColor(0xFD, 0xEC, 0xEB)
FILL_GRAY = RGBColor(0xF3, 0xF6, 0xF9)
FILL_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def remove_all_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def render_svg_with_chrome(svg_path: Path, png_path: Path) -> None:
    if not CHROME_PATH.exists():
        raise FileNotFoundError(f"Google Chrome not found: {CHROME_PATH}")

    png_path.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix="bogp_chrome_profile_") as profile_dir:
        process = subprocess.Popen(
            [
                str(CHROME_PATH),
                "--headless=new",
                "--disable-gpu",
                "--hide-scrollbars",
                "--no-first-run",
                "--no-default-browser-check",
                "--disable-background-networking",
                "--disable-sync",
                "--disable-default-apps",
                "--mute-audio",
                f"--user-data-dir={profile_dir}",
                "--window-size=2400,1350",
                f"--screenshot={png_path}",
                svg_path.resolve().as_uri(),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )

        deadline = time.monotonic() + 30
        last_size = -1
        stable_polls = 0
        while time.monotonic() < deadline:
            if png_path.exists():
                size = png_path.stat().st_size
                if size > 0 and size == last_size:
                    stable_polls += 1
                    if stable_polls >= 2:
                        break
                else:
                    last_size = size
                    stable_polls = 0

            if process.poll() is not None and png_path.exists() and png_path.stat().st_size > 0:
                break
            time.sleep(0.25)

        if process.poll() is None:
            process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)

        stdout, stderr = process.communicate()
        if not png_path.exists() or png_path.stat().st_size == 0:
            raise RuntimeError(
                f"Failed to render {svg_path.name} with Chrome.\nstdout:\n{stdout}\nstderr:\n{stderr}"
            )


def svg_to_png_assets() -> dict[str, Path]:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    png_paths: dict[str, Path] = {}
    for key, svg_path in FIGURE_MAP.items():
        png_path = GENERATED_DIR / f"{svg_path.stem}.png"
        render_svg_with_chrome(svg_path, png_path)
        png_paths[key] = png_path
    return png_paths


def set_placeholder_text(placeholder, text: str, size_pt: float | None = None, align=PP_ALIGN.LEFT) -> None:
    placeholder.text = text
    tf = placeholder.text_frame
    tf.word_wrap = True
    for paragraph in tf.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.color.rgb = COLOR_TITLE
            if size_pt is not None:
                run.font.size = Pt(size_pt)


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    paragraphs,
    font_size: float = 20,
    color: RGBColor = COLOR_BODY,
    bold: bool = False,
    fill_color: RGBColor | None = None,
    line_color: RGBColor | None = None,
    rounded: bool = False,
    align=PP_ALIGN.LEFT,
    margin: float = 0.08,
):
    if fill_color is not None or line_color is not None or rounded:
        shape_type = (
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
            if rounded
            else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        )
        shape = slide.shapes.add_shape(
            shape_type,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        if fill_color is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color
        else:
            shape.fill.background()

        if line_color is not None:
            shape.line.color.rgb = line_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
    else:
        shape = slide.shapes.add_textbox(
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )

    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)

    for idx, item in enumerate(paragraphs):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, str):
            item = {"text": item}
        paragraph.text = item["text"]
        paragraph.alignment = item.get("align", align)
        paragraph.space_before = Pt(item.get("space_before", 0))
        paragraph.space_after = Pt(item.get("space_after", 3))
        for run in paragraph.runs:
            run.font.size = Pt(item.get("size", font_size))
            run.font.bold = item.get("bold", bold)
            run.font.color.rgb = item.get("color", color)
    return shape


def add_panel(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
    body_lines: list[str],
    fill_color: RGBColor,
    line_color: RGBColor,
    title_size: float = 22,
    body_size: float = 18,
    center: bool = False,
):
    paragraphs = [
        {"text": title, "size": title_size, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER if center else PP_ALIGN.LEFT},
    ]
    for line in body_lines:
        paragraphs.append(
            {
                "text": line,
                "size": body_size,
                "color": COLOR_BODY,
                "align": PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
                "space_after": 2,
            }
        )
    return add_textbox(
        slide,
        left,
        top,
        width,
        height,
        paragraphs,
        fill_color=fill_color,
        line_color=line_color,
        rounded=True,
        align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,
    )


def add_picture(slide, image_path: Path, left: float, top: float, width: float | None = None, height: float | None = None):
    kwargs = {}
    if width is not None:
        kwargs["width"] = Inches(width)
    if height is not None:
        kwargs["height"] = Inches(height)
    return slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), **kwargs)


def add_picture_contain(
    slide,
    image_path: Path,
    left: float,
    top: float,
    max_width: float,
    max_height: float,
    center_x: bool = True,
    center_y: bool = False,
):
    with Image.open(image_path) as image:
        image_width, image_height = image.size

    scale = min(max_width / image_width, max_height / image_height)
    width = image_width * scale
    height = image_height * scale

    adjusted_left = left + (max_width - width) / 2 if center_x else left
    adjusted_top = top + (max_height - height) / 2 if center_y else top
    return add_picture(slide, image_path, adjusted_left, adjusted_top, width=width, height=height)


def add_connector(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = COLOR_MUTED, width_pt: float = 2.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width_pt)
    return line


def add_chevron_row(slide, items, left: float, top: float, width: float, height: float):
    gap = 0.1
    item_width = (width - gap * (len(items) - 1)) / len(items)
    for idx, item in enumerate(items):
        x = left + idx * (item_width + gap)
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.CHEVRON,
            Inches(x),
            Inches(top),
            Inches(item_width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = item["fill"]
        shape.line.color.rgb = item["line"]
        shape.line.width = Pt(1.5)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        entries = [
            {"text": item["title"], "size": 20, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER},
            {"text": item["body"], "size": 14, "color": COLOR_BODY, "align": PP_ALIGN.CENTER},
        ]
        for p_idx, entry in enumerate(entries):
            p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()
            p.text = entry["text"]
            p.alignment = entry["align"]
            p.space_after = Pt(1)
            for run in p.runs:
                run.font.size = Pt(entry["size"])
                run.font.bold = entry.get("bold", False)
                run.font.color.rgb = entry["color"]


def make_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_placeholder_text(slide.shapes.title, "交叉率・突然変異率をベイズ最適化で\n動的に調整する遺伝的プログラミング", 28, align=PP_ALIGN.CENTER)
    subtitle = slide.placeholders[1]
    set_placeholder_text(subtitle, "進捗報告", 22, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        4.6,
        4.35,
        4.1,
        0.55,
        [{"text": "2026.04.12", "size": 24, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        3.85,
        5.52,
        5.6,
        0.45,
        [{"text": "閉ループ制御・目的関数設計・今後の実装方針", "size": 17, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )


def make_agenda_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "本日の報告内容")
    add_textbox(
        slide,
        0.92,
        1.08,
        6.8,
        0.5,
        [{"text": "研究ノート 1-8 章をもとに、設計の芯を中心に整理する。", "size": 20, "color": COLOR_MUTED}],
    )

    items = [
        ("1", "背景と目的", FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8)),
        ("2", "研究課題", FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8)),
        ("3", "閉ループ制御", FILL_BLUE, COLOR_BLUE),
        ("4", "目的関数とBO", FILL_GREEN, COLOR_GREEN),
        ("5", "今後の方針", FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8)),
    ]
    for idx, (num, text, fill, line) in enumerate(items):
        left = 0.92 + idx * 2.45
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left), Inches(2.05), Inches(0.55), Inches(0.55))
        circle.fill.solid()
        circle.fill.fore_color.rgb = line
        circle.line.fill.background()
        tf = circle.text_frame
        tf.text = num
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        for run in tf.paragraphs[0].runs:
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = FILL_WHITE
        add_panel(slide, left + 0.18, 2.48, 2.1, 1.15, text, [], fill, line, title_size=18, body_size=16, center=True)
    add_textbox(
        slide,
        2.0,
        4.55,
        9.3,
        1.0,
        [
            {"text": "重点", "size": 20, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER},
            {"text": "Slide 7-15 で制御構造、文脈設計、目的関数、文脈付き BO 採用理由を詳しく説明する。", "size": 18, "color": COLOR_BODY, "align": PP_ALIGN.CENTER},
        ],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def make_theme_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "研究テーマと基本アイデア")
    add_panel(
        slide,
        0.85,
        1.25,
        4.15,
        2.05,
        "研究テーマ",
        [
            "・交叉率 p_c と突然変異率 p_m を動的制御する",
            "・上位層に BO を置き、進化状態に応じて率を更新する",
            "・GP を制御対象、BO を制御器として扱う",
        ],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=24,
        body_size=18,
    )
    add_panel(
        slide,
        0.85,
        3.62,
        4.15,
        1.2,
        "狙い",
        ["進化計算をオンラインで閉ループ制御する枠組みへ拡張する。"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=24,
        body_size=18,
    )

    add_panel(slide, 5.6, 2.2, 2.1, 1.3, "GP状態", ["現在の進化状態"], FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8), center=True)
    add_panel(slide, 8.05, 2.2, 2.1, 1.3, "BO", ["次の率を提案"], FILL_BLUE, COLOR_BLUE, center=True)
    add_panel(slide, 10.5, 2.2, 2.1, 1.3, "操作率", ["p_c, p_m"], FILL_GREEN, COLOR_GREEN, center=True)
    add_connector(slide, 7.7, 2.85, 8.05, 2.85)
    add_connector(slide, 10.15, 2.85, 10.5, 2.85)
    add_textbox(
        slide,
        6.0,
        4.25,
        6.1,
        1.2,
        [
            {"text": "考え方", "size": 22, "bold": True, "color": COLOR_TITLE},
            {"text": "固定値のチューニングではなく、進行中の GP に対する逐次制御として扱う。", "size": 20, "color": COLOR_BODY},
        ],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
    )


def make_background_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "背景: なぜ固定率では不十分か")
    add_textbox(
        slide,
        0.9,
        1.12,
        5.4,
        1.6,
        [
            {"text": "・GP では操作率が探索挙動を大きく左右する", "size": 21},
            {"text": "・望ましい率は進化段階で変化する", "size": 21},
            {"text": "・固定率では探索性と収束性の両立が難しい", "size": 21},
        ],
    )
    add_chevron_row(
        slide,
        [
            {"title": "探索期", "body": "突然変異率を高めて\n多様性を確保", "fill": FILL_AMBER, "line": COLOR_AMBER},
            {"title": "収束移行", "body": "交叉率を高めて\n有望解を組み替え", "fill": FILL_BLUE, "line": COLOR_BLUE},
            {"title": "停滞打破", "body": "探索性を戻して\n再探索を促す", "fill": FILL_GREEN, "line": COLOR_GREEN},
        ],
        0.92,
        3.0,
        11.35,
        1.55,
    )
    add_textbox(
        slide,
        1.8,
        5.1,
        9.6,
        1.0,
        [{"text": "同じ率を全世代に適用するより、状態に応じて率を変える方が自然である。", "size": 22, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def make_objective_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "本研究の目的")
    add_panel(
        slide,
        0.92,
        1.25,
        4.05,
        2.0,
        "目的",
        [
            "・操作率を状態依存の制御入力として扱う",
            "・収束性能、多様性、収束時間を改善する",
        ],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=24,
        body_size=18,
    )
    add_panel(
        slide,
        0.92,
        3.55,
        4.05,
        1.35,
        "比較",
        ["固定率 GP と手設計スケジュールを基準にする"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=24,
        body_size=18,
    )
    cards = [
        ("固定率 GP", FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8)),
        ("手設計\nスケジュール", FILL_AMBER, COLOR_AMBER),
        ("提案法\n文脈付き BO", FILL_BLUE, COLOR_BLUE),
    ]
    for idx, (title, fill, line) in enumerate(cards):
        add_panel(slide, 5.6 + idx * 2.35, 2.2, 2.1, 1.7, title, [], fill, line, title_size=20, center=True)
    add_textbox(
        slide,
        6.0,
        4.45,
        5.9,
        0.9,
        [{"text": "安定して良い進化を行える制御則の設計を目指す。", "size": 21, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def make_challenge_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "研究課題: 最適率は状態に依存する")
    add_textbox(
        slide,
        0.95,
        1.1,
        11.2,
        0.65,
        [{"text": "同じ (p_c, p_m) でも、探索初期と終盤では効果が異なる。", "size": 22, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )
    add_panel(
        slide,
        0.92,
        1.95,
        5.45,
        2.65,
        "単純な最適化では不十分",
        [
            "・f(p_c, p_m) だけを学習すると時変性を扱えない",
            "・探索初期向きの率と終盤向きの率が混ざる",
            "・BO から見ると目的関数が非定常になる",
        ],
        FILL_RED,
        COLOR_RED,
        title_size=24,
        body_size=18,
    )
    add_panel(
        slide,
        6.96,
        1.95,
        5.45,
        2.65,
        "採用する考え方",
        [
            "・現在状態 c_t を文脈として BO に入力する",
            "・r_t = g(c_t, p_c, p_m) + ε として報酬を学習する",
            "・状態に応じて次の率を選ぶ",
        ],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=24,
        body_size=18,
    )
    add_textbox(
        slide,
        3.15,
        5.05,
        7.2,
        0.95,
        [{"text": "文脈付き BO にすることで、率の良し悪しを「今の状態」に結び付けて扱える。", "size": 21, "color": COLOR_TITLE, "bold": True, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def make_figure_slide_7(prs: Presentation, png_assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "提案手法: GP と BO の閉ループ構造")
    add_textbox(
        slide,
        0.92,
        1.04,
        5.5,
        0.85,
        [
            {"text": "・GP をプラント、BO を制御器として扱う", "size": 19},
            {"text": "・制御入力: u_t = [p_c, p_m]", "size": 19},
            {"text": "・観測量: HV、多様性、停滞、bloat", "size": 19},
        ],
        fill_color=FILL_WHITE,
        line_color=RGBColor(0xD8, 0xE0, 0xE8),
        rounded=True,
    )
    add_picture_contain(slide, png_assets["slide07"], 1.05, 1.98, max_width=11.15, max_height=4.95)


def make_figure_slide_8(prs: Presentation, png_assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "1 制御ステップで何をしているか")
    add_textbox(
        slide,
        1.15,
        1.0,
        10.8,
        0.42,
        [{"text": "観測 → 提案 → 実行 → 評価 → 更新 を 1 サイクルとして繰り返す。", "size": 19, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )
    add_picture_contain(slide, png_assets["slide08"], 1.15, 1.44, max_width=10.95, max_height=5.75)


def make_interval_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "なぜ数世代ごとの区間制御にするのか")
    add_panel(
        slide,
        0.92,
        1.2,
        3.0,
        1.45,
        "問題点",
        ["・1 世代評価はノイズが大きい", "・BO の観測が不安定になる"],
        FILL_RED,
        COLOR_RED,
        title_size=22,
        body_size=18,
    )
    add_panel(
        slide,
        4.18,
        1.2,
        3.05,
        1.45,
        "採用案",
        ["・制御区間を 3-5 世代にする", "・区間中は率を固定する"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=22,
        body_size=18,
    )
    add_panel(
        slide,
        7.5,
        1.2,
        4.0,
        1.45,
        "効果",
        ["・応答性を残しつつノイズを低減", "・区間全体の改善量で評価しやすい"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=22,
        body_size=18,
    )
    add_textbox(
        slide,
        0.98,
        3.25,
        11.3,
        0.6,
        [{"text": "制御区間のイメージ", "size": 20, "bold": True, "color": COLOR_TITLE}],
    )
    x0 = 1.1
    box_w = 0.78
    for idx in range(9):
        left = x0 + idx * 1.08
        fill = FILL_BLUE if idx < 3 else FILL_GREEN if idx < 6 else FILL_AMBER
        line = COLOR_BLUE if idx < 3 else COLOR_GREEN if idx < 6 else COLOR_AMBER
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(4.0), Inches(box_w), Inches(0.72))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = line
        tf = shape.text_frame
        tf.text = str(idx + 1)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        for run in tf.paragraphs[0].runs:
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = COLOR_TITLE
    add_textbox(slide, 1.4, 4.92, 2.2, 0.45, [{"text": "区間 1", "size": 18, "bold": True, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_textbox(slide, 4.62, 4.92, 2.2, 0.45, [{"text": "区間 2", "size": 18, "bold": True, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_textbox(slide, 7.86, 4.92, 2.2, 0.45, [{"text": "区間 3", "size": 18, "bold": True, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        2.2,
        5.62,
        8.7,
        0.7,
        [{"text": "各区間の最後に報酬を集計し、次の p_c と p_m を更新する。", "size": 20, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def make_figure_slide_10(prs: Presentation, png_assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "文脈ベクトル c_t の設計")
    add_panel(
        slide,
        0.85,
        1.3,
        3.65,
        3.15,
        "文脈要素",
        [
            "・tau: 進化段階",
            "・HV, DeltaHV: 性能と改善速度",
            "・D: 探索性",
            "・s, b: 停滞と複雑化",
        ],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=23,
        body_size=18,
    )
    add_textbox(
        slide,
        0.95,
        4.82,
        3.45,
        0.9,
        [{"text": "何を BO に見せるかで、次の率提案の質が決まる。", "size": 19, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )
    add_picture_contain(slide, png_assets["slide10"], 4.78, 1.25, max_width=7.95, max_height=4.85)


def make_constraints_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "操作率にどのような制約を入れるか")
    add_panel(
        slide,
        0.92,
        1.28,
        3.45,
        1.8,
        "範囲制約",
        ["・p_c ∈ [0.55, 0.95]", "・p_m ∈ [0.01, 0.30]"],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=23,
        body_size=18,
    )
    add_panel(
        slide,
        4.62,
        1.28,
        3.2,
        1.8,
        "組合せ制約",
        ["・必要なら p_c + p_m ≤ 1.0", "・極端な率の組合せを防ぐ"],
        FILL_GREEN,
        COLOR_GREEN,
        title_size=23,
        body_size=18,
    )
    add_panel(
        slide,
        8.08,
        1.28,
        4.0,
        1.8,
        "変化量制約",
        ["・|Δp_c|, |Δp_m| ≤ Δmax", "・急激な挙動変化を防ぐ"],
        FILL_AMBER,
        COLOR_AMBER,
        title_size=23,
        body_size=18,
    )
    add_textbox(
        slide,
        1.1,
        3.8,
        10.7,
        1.35,
        [
            {"text": "制約を入れる理由", "size": 22, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER},
            {"text": "BO の自由度を残しつつ、非現実的な提案や探索の破綻を防ぐため。", "size": 20, "align": PP_ALIGN.CENTER},
        ],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        2.3,
        5.45,
        8.2,
        0.55,
        [{"text": "安全策としての制約であり、探索性そのものを失わせるためのものではない。", "size": 18, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )


def make_figure_slide_12(prs: Presentation, png_assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "目的関数: 何を良い制御とみなすか")
    add_textbox(
        slide,
        1.05,
        1.05,
        10.8,
        0.7,
        [
            {"text": "HV 改善を主軸にしつつ、多様性維持、停滞回避、bloat 抑制を同時に評価する。", "size": 19, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER},
        ],
        fill_color=FILL_WHITE,
        line_color=RGBColor(0xD8, 0xE0, 0xE8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )
    add_picture_contain(slide, png_assets["slide12"], 1.1, 1.83, max_width=11.0, max_height=5.15)


def make_reason_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "なぜこの目的関数にしたのか")
    add_panel(
        slide,
        0.92,
        1.25,
        4.2,
        2.2,
        "設計意図",
        [
            "・多様性は単純最大化しない",
            "・目標多様性 D*(tau) に追従させる",
            "・初期は探索、後半は収束を許容する",
        ],
        FILL_BLUE,
        COLOR_BLUE,
        title_size=23,
        body_size=18,
    )
    add_textbox(
        slide,
        5.55,
        1.42,
        6.1,
        0.52,
        [{"text": "目標多様性 D*(tau)", "size": 21, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )
    add_connector(slide, 6.0, 4.7, 10.9, 4.7, COLOR_MUTED, 2.5)
    add_connector(slide, 6.0, 4.7, 6.0, 2.1, COLOR_MUTED, 2.5)
    add_connector(slide, 6.3, 2.45, 10.45, 4.25, COLOR_BLUE, 4.0)
    add_textbox(slide, 5.78, 2.0, 0.6, 0.35, [{"text": "高", "size": 16, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_textbox(slide, 5.75, 4.72, 0.65, 0.35, [{"text": "0", "size": 16, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_textbox(slide, 10.7, 4.72, 0.65, 0.35, [{"text": "tau", "size": 16, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_textbox(slide, 10.15, 4.05, 0.8, 0.35, [{"text": "低", "size": 16, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
    add_panel(
        slide,
        5.65,
        5.15,
        5.9,
        0.9,
        "初期重み",
        ["w_hv=0.55   w_d=0.25   w_s=0.10   w_b=0.10"],
        FILL_GRAY,
        RGBColor(0x8A, 0x98, 0xA8),
        title_size=20,
        body_size=18,
        center=True,
    )


def make_bo_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "BO 側はどのように次の率を決めるか")
    add_textbox(
        slide,
        0.92,
        1.15,
        10.9,
        0.8,
        [
            {"text": "・回帰器: Gaussian Process Regressor", "size": 20},
            {"text": "・取得関数: Expected Improvement", "size": 20},
            {"text": "・文脈固定のもとで次の (p_c, p_m) を選ぶ", "size": 20},
        ],
    )
    add_panel(slide, 1.0, 3.0, 2.6, 1.25, "観測データ", ["D_t = {(c_i, u_i, r_i)}"], FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8), center=True)
    add_panel(slide, 4.15, 3.0, 2.8, 1.25, "GPR", ["Matern + White noise"], FILL_BLUE, COLOR_BLUE, center=True)
    add_panel(slide, 7.5, 3.0, 2.3, 1.25, "EI", ["改善期待値を評価"], FILL_GREEN, COLOR_GREEN, center=True)
    add_panel(slide, 10.25, 3.0, 1.85, 1.25, "次の率", ["(p_c, p_m)"], FILL_AMBER, COLOR_AMBER, center=True)
    add_connector(slide, 3.6, 3.62, 4.15, 3.62)
    add_connector(slide, 6.95, 3.62, 7.5, 3.62)
    add_connector(slide, 9.8, 3.62, 10.25, 3.62)
    add_textbox(
        slide,
        2.25,
        5.0,
        8.8,
        0.95,
        [{"text": "観測が少ない段階でも、不確実性を含めて次の率を提案できることが BO の利点である。", "size": 20, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_GRAY,
        line_color=RGBColor(0x8A, 0x98, 0xA8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def make_figure_slide_15(prs: Presentation, png_assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "なぜ文脈付き BO を採用するのか")
    add_textbox(
        slide,
        1.1,
        0.98,
        10.8,
        0.48,
        [{"text": "同じ率でも、状態が違えば望ましい報酬は変わる。", "size": 19, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )
    add_picture_contain(slide, png_assets["slide15"], 1.05, 1.55, max_width=11.1, max_height=5.35)


def make_stages_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "実装の段階分けと評価方法")
    add_panel(slide, 0.92, 1.28, 3.7, 1.1, "段階 1", ["簡易シミュレータで制御器単体を確認"], FILL_BLUE, COLOR_BLUE, title_size=22, body_size=18)
    add_panel(slide, 0.92, 2.58, 3.7, 1.1, "段階 2", ["実 GP へ接続して指標を測定"], FILL_GREEN, COLOR_GREEN, title_size=22, body_size=18)
    add_panel(slide, 0.92, 3.88, 3.7, 1.1, "段階 3", ["固定率法・文脈なし BO と比較"], FILL_AMBER, COLOR_AMBER, title_size=22, body_size=18)
    add_textbox(
        slide,
        5.1,
        1.22,
        6.8,
        0.55,
        [{"text": "評価指標", "size": 23, "bold": True, "color": COLOR_TITLE}],
    )
    metrics = [
        ("最終 HV", FILL_BLUE, COLOR_BLUE),
        ("平均多様性", FILL_GREEN, COLOR_GREEN),
        ("HV 到達世代数", FILL_AMBER, COLOR_AMBER),
        ("停滞頻度", FILL_RED, COLOR_RED),
        ("平均木サイズ", FILL_GRAY, RGBColor(0x8A, 0x98, 0xA8)),
    ]
    for idx, (text, fill, line) in enumerate(metrics):
        row = idx // 2
        col = idx % 2
        if idx == 4:
            left = 6.6
            top = 4.6
        else:
            left = 5.2 + col * 3.3
            top = 1.9 + row * 1.3
        add_panel(slide, left, top, 2.8, 0.92, text, [], fill, line, title_size=20, center=True)


def make_conditions_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "現時点での実験条件")
    rows = [
        ("実装言語", "Python"),
        ("GP ライブラリ候補", "DEAP"),
        ("BO 実装", "scikit-learn ベース"),
        ("制御区間", "3 世代"),
        ("初期設計点数", "8 点"),
        ("候補評価点数", "512 点"),
    ]
    left = 2.1
    top = 1.55
    row_h = 0.65
    for idx, (label, value) in enumerate(rows):
        y = top + idx * row_h
        add_textbox(
            slide,
            left,
            y,
            3.1,
            0.56,
            [{"text": label, "size": 20, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
            fill_color=FILL_GRAY,
            line_color=RGBColor(0x8A, 0x98, 0xA8),
            rounded=False,
            align=PP_ALIGN.CENTER,
        )
        add_textbox(
            slide,
            left + 3.1,
            y,
            4.9,
            0.56,
            [{"text": value, "size": 20, "color": COLOR_BODY, "align": PP_ALIGN.CENTER}],
            fill_color=FILL_WHITE,
            line_color=RGBColor(0xD8, 0xE0, 0xE8),
            rounded=False,
            align=PP_ALIGN.CENTER,
        )
    add_textbox(
        slide,
        2.2,
        5.9,
        7.8,
        0.52,
        [{"text": "まずは研究仮説を確認するための最小構成で実験を開始する。", "size": 18, "color": COLOR_MUTED, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
    )


def make_summary_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    set_placeholder_text(slide.shapes.title, "現時点の設計方針まとめ")
    add_panel(slide, 0.92, 1.35, 5.25, 1.15, "方針 1", ["数世代ごとの区間制御"], FILL_BLUE, COLOR_BLUE, title_size=21, body_size=19)
    add_panel(slide, 0.92, 2.72, 5.25, 1.15, "方針 2", ["文脈付き BO による状態依存の率調整"], FILL_GREEN, COLOR_GREEN, title_size=21, body_size=19)
    add_panel(slide, 0.92, 4.09, 5.25, 1.15, "方針 3", ["HV 改善量と多様性維持の合成報酬"], FILL_AMBER, COLOR_AMBER, title_size=21, body_size=19)
    add_panel(slide, 0.92, 5.46, 5.25, 1.15, "方針 4", ["停滞と bloat を明示的に抑制"], FILL_RED, COLOR_RED, title_size=21, body_size=19)

    add_panel(
        slide,
        7.0,
        2.1,
        4.5,
        1.6,
        "次の課題",
        ["・実 GP 実装へ接続", "・実問題で HV を測定", "・固定率法や文脈なし BO と比較"],
        FILL_GRAY,
        RGBColor(0x8A, 0x98, 0xA8),
        title_size=24,
        body_size=18,
    )
    add_connector(slide, 8.2, 4.55, 10.3, 4.55, COLOR_BLUE, 3.0)
    add_textbox(
        slide,
        7.3,
        4.95,
        3.9,
        0.62,
        [{"text": "実 GP 接続 → 比較実験", "size": 22, "bold": True, "color": COLOR_TITLE, "align": PP_ALIGN.CENTER}],
        fill_color=FILL_WHITE,
        line_color=RGBColor(0xD8, 0xE0, 0xE8),
        rounded=True,
        align=PP_ALIGN.CENTER,
    )


def build_presentation() -> Path:
    png_assets = svg_to_png_assets()
    prs = Presentation(str(TEMPLATE_PATH))
    remove_all_slides(prs)
    prs.core_properties.title = "260412_A"
    prs.core_properties.subject = "Bayesian optimization controlled GP progress report"

    make_title_slide(prs)
    make_agenda_slide(prs)
    make_theme_slide(prs)
    make_background_slide(prs)
    make_objective_slide(prs)
    make_challenge_slide(prs)
    make_figure_slide_7(prs, png_assets)
    make_figure_slide_8(prs, png_assets)
    make_interval_slide(prs)
    make_figure_slide_10(prs, png_assets)
    make_constraints_slide(prs)
    make_figure_slide_12(prs, png_assets)
    make_reason_slide(prs)
    make_bo_slide(prs)
    make_figure_slide_15(prs, png_assets)
    make_stages_slide(prs)
    make_conditions_slide(prs)
    make_summary_slide(prs)

    prs.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_presentation()
    print(path)
